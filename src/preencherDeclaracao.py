from docx import Document
from datetime import datetime
import subprocess
import os

# ===================== CONFIGURAÇÃO MANUAL =====================
CHAMADO = "SRV-1181784"
TELEFONE = "81 6000-0527"
CAMINHO_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
# ===============================================================

texto_bruto = """
Nome: Nome
Usuário: Usuario
 
Equipamento despachado via Sedex no dia 00/00/0000
Rua da Manga, 123
Bairro: Mangueira
Cidade: Uberlândia - MG
CEP: 00000-000
 
Dados do Equipamento enviado:
Notebook Modelo: Lenovo Thinkpad L14
Patrimônio: PTR-123456
Serial: 123456
"""


def definir_valor_por_modelo(modelo):
    modelo = modelo.upper()
    if "L14" in modelo or "E14" in modelo:
        return "R$ 5.200,00"
    if "3420" in modelo or "3410" in modelo:
        return "R$ 4.200,00"
    return "R$ 0,00"


def extrair_dados(texto):
    dados = {}
    nome_dest = ""

    linhas = list(dict.fromkeys(
        [l.strip() for l in texto.splitlines() if l.strip()]
    ))

    rua = bairro = cidade = cep = ""
    modelo_envio = patrimonio_envio = ""
    lendo_envio = False

    for linha in linhas:
        lower = linha.lower()

        if lower.startswith("nome:"):
            nome_dest = linha.split(":", 1)[1].strip()
            dados["{{nome_dest}}"] = nome_dest

        elif lower.startswith("bairro:"):
            bairro = linha.split(":", 1)[1].strip()

        elif lower.startswith("cidade:"):
            cidade = linha.split(":", 1)[1].strip()

        elif lower.startswith("cep:"):
            cep = linha.split(":", 1)[1].strip()

        elif lower.startswith("rua") or lower.startswith("av"):
            rua = linha

        elif "dados do equipamento enviado" in lower:
            lendo_envio = True

        elif lendo_envio and lower.startswith("notebook modelo:"):
            modelo_envio = linha.split(":", 1)[1].strip()

        elif lendo_envio and lower.startswith("patrimônio:"):
            patrimonio_envio = linha.split(":", 1)[1].strip()

    dados["{{endereco}}"] = (
        f"{rua}\n"
        f"Bairro: {bairro}\n"
        f"Cidade: {cidade}"
    )

    dados["{{cidade}}"] = cidade
    dados["{{cep}}"] = cep
    dados["{{telefone}}"] = TELEFONE
    dados["{{marca_modelo}}"] = modelo_envio
    dados["{{patrimonio}}"] = patrimonio_envio
    dados["{{chamado}}"] = CHAMADO
    dados["{{valor}}"] = definir_valor_por_modelo(modelo_envio)

    hoje = datetime.now()
    dados["{{dia_data}}"] = hoje.strftime("%d")
    dados["{{mes_data}}"] = hoje.strftime("%m")

    return dados, nome_dest


def substituir_documento(doc, dados):
    def substituir(p):
        texto = p.text
        for k, v in dados.items():
            texto = texto.replace(k, v)
        if texto != p.text:
            p.clear()
            p.add_run(texto)

    for p in doc.paragraphs:
        substituir(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    substituir(p)


def limpar_placeholders(doc):
    for p in doc.paragraphs:
        if "{{" in p.text:
            p.clear()

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{" in p.text:
                        p.clear()


def converter_para_pdf(docx_path):
    subprocess.run(
        [
            CAMINHO_SOFFICE,
            "--headless",
            "--convert-to", "pdf",
            docx_path,
            "--outdir", os.path.dirname(docx_path) or "."
        ],
        check=True
    )


# ===================== EXECUÇÃO =====================
dados, nome_dest = extrair_dados(texto_bruto)

docx_saida = f"DECLARAÇÃO DE CONTEÚDO {nome_dest}.docx"

doc = Document("geradorDeclaracaoConteudo/template/Declaracao.docx")
substituir_documento(doc, dados)
limpar_placeholders(doc)
doc.save(docx_saida)

converter_para_pdf(docx_saida)

print("✅ DOCX e PDF gerados corretamente!")
